import hashlib
import json

import pytest
from fastapi import BackgroundTasks

from app.api.tasks import (
    ToolDecision,
    decide_tool_call_endpoint,
    list_task_artifacts_endpoint,
)
from app.db.database import init_db
from app.models.artifact_repository import get_artifact, list_artifacts
from app.models.conversation_repository import create_conversation
from app.models.project_repository import create_project
from app.models.repository import create_model
from app.models.task_event_repository import list_task_events
from app.models.task_repository import create_task, get_task
from app.models.tool_call_repository import create_tool_call, get_tool_call, list_tool_calls
from app.models.tool_result_repository import list_tool_results
from app.services.agent import run_agent
from app.services.tools import approve_tool_call, execute_tool_call


@pytest.fixture(autouse=True)
def test_database(monkeypatch, tmp_path):
    monkeypatch.setattr("app.db.database.DATABASE_PATH", tmp_path / "artifacts.db")
    init_db()


def _task_context(tmp_path, name, goal="Create an artifact"):
    workspace = tmp_path / name.replace(" ", "_")
    workspace.mkdir()
    project = create_project(name, str(workspace))
    model = create_model("Selected", "ollama", "selected:test")
    conversation = create_conversation(name, model.id, project_id=project.id)
    task = create_task(
        type="document",
        mode="agent",
        conversation_id=conversation.id,
        project_id=project.id,
        plan={"goal": goal},
    )
    return task, conversation, workspace


@pytest.mark.parametrize(
    "fmt,path",
    [
        ("markdown", "Report.md"),
        ("pdf", "Report.pdf"),
        ("docx", "Report.docx"),
        ("xlsx", "Report.xlsx"),
        ("pptx", "Report.pptx"),
    ],
)
def test_artifact_tool_generates_reopens_and_persists_each_format(tmp_path, fmt, path):
    task, conversation, workspace = _task_context(tmp_path, f"Artifact {fmt}")
    markdown = (
        "# Mosaic Artifact Report\n\n"
        "This artifact must reopen and preserve its expected content.\n\n"
        "## Evidence\n- Validation passed.\n"
    )
    content = (
        json.dumps({
            "sheets": [{
                "name": "Data",
                "rows": [
                    ["Name", "Value"],
                    ["Mosaic", 42],
                    ["=not a formula", {"formula": "=SUM(B2:B2)"}],
                ],
            }]
        })
        if fmt == "xlsx" else markdown
    )
    call = create_tool_call(
        conversation.id,
        "artifact",
        {"format": fmt, "path": path, "content": content, "title": "Mosaic Artifact Report"},
        task_id=task.id,
    )
    result = execute_tool_call(approve_tool_call(call))

    assert result.status == "completed"
    assert result.data["format"] == fmt
    assert result.data["reopened"] is True
    assert result.data["validation"]["valid"] is True
    assert result.data["sha256"] == hashlib.sha256((workspace / path).read_bytes()).hexdigest()
    assert get_tool_call(call.id).status == "completed"
    assert len(list_tool_results(call.id)) == 1

    artifact = get_artifact(result.data["artifact_id"])
    assert artifact is not None
    assert artifact.task_id == task.id
    assert artifact.type == fmt
    assert artifact.path == path
    assert artifact.metadata["sha256"] == result.data["sha256"]
    assert artifact.metadata["validation"]["valid"] is True
    assert list_artifacts(task.id) == [artifact]

    if fmt == "markdown":
        assert (workspace / path).read_text(encoding="utf-8") == markdown
    elif fmt == "pdf":
        from pypdf import PdfReader
        reopened = PdfReader(str(workspace / path))
        assert len(reopened.pages) >= 1
        assert "Mosaic Artifact Report" in "\n".join(page.extract_text() or "" for page in reopened.pages)
    elif fmt == "docx":
        from docx import Document
        reopened = Document(str(workspace / path))
        assert "Mosaic Artifact Report" in "\n".join(p.text for p in reopened.paragraphs)
    elif fmt == "xlsx":
        from openpyxl import load_workbook
        reopened = load_workbook(workspace / path, data_only=False)
        assert reopened.sheetnames == ["Data"]
        assert reopened["Data"]["B2"].value == 42
        assert reopened["Data"]["A3"].value == "'=not a formula"
        assert reopened["Data"]["B3"].value == "=SUM(B2:B2)"
    else:
        from pptx import Presentation
        reopened = Presentation(str(workspace / path))
        text = "\n".join(
            shape.text for slide in reopened.slides for shape in slide.shapes if hasattr(shape, "text")
        )
        assert len(reopened.slides) == 2
        assert "Mosaic Artifact Report" in text
        assert "Validation passed." in text


def test_artifact_listing_invalidates_evidence_when_file_bytes_change(tmp_path):
    task, conversation, workspace = _task_context(tmp_path, "Artifact invalidation")
    call = create_tool_call(
        conversation.id,
        "artifact",
        {
            "format": "pdf",
            "path": "Report.pdf",
            "content": "# Stable Report\n\nOriginal verified content.",
            "title": "Stable Report",
        },
        task_id=task.id,
    )
    result = execute_tool_call(approve_tool_call(call))
    artifact = get_artifact(result.data["artifact_id"])
    assert artifact is not None

    (workspace / "Report.pdf").write_bytes(b"tampered bytes")
    listing = list_task_artifacts_endpoint(task.id)
    item = listing["items"][0]
    assert item["current_verification"] == "invalidated"
    assert item["current_sha256"] != artifact.metadata["sha256"]
    events = list_task_events(task.id)
    assert any(
        event.event_name == "artifact.evidence_invalidated"
        and event.metadata.get("artifact_id") == artifact.id
        for event in events
    )


def test_pdf_task_flow_requires_approval_then_completes_without_second_model_call(monkeypatch, tmp_path):
    goal = "Create a PDF report named Report.pdf"
    task, conversation, workspace = _task_context(tmp_path, "PDF Task Flow", goal)
    markdown = "# Report\n\nA report created through the actual Task and ToolCall lifecycle."

    result = run_agent(task.id, lambda context: {"type": "final", "content": markdown})
    assert result["status"] == "awaiting_approval"
    pending = [
        call for call in list_tool_calls(conversation.id)
        if call.task_id == task.id and call.tool_name == "artifact"
    ]
    assert len(pending) == 1
    assert get_task(task.id).status == "awaiting_approval"

    monkeypatch.setattr(
        "app.api.tasks.create_provider_decider",
        lambda _task_id: lambda _context: (_ for _ in ()).throw(AssertionError("unexpected second model call")),
    )
    background = BackgroundTasks()
    resumed = decide_tool_call_endpoint(
        task.id, pending[0].id, ToolDecision(decision="approve"), background
    )
    assert resumed.status == "executing"
    assert len(background.tasks) == 1
    scheduled = background.tasks[0]
    scheduled.func(*scheduled.args, **scheduled.kwargs)

    completed = get_task(task.id)
    assert completed.status == "completed", completed.result
    assert completed.result["artifact"]["path"] == "Report.pdf"
    assert completed.result["artifact"]["artifact_id"] > 0
    assert completed.result["artifact"]["verification_event_recorded"] is True
    assert (workspace / "Report.pdf").is_file()
    assert list_artifacts(task.id)[0].metadata["sha256"] == completed.result["artifact"]["sha256"]
    assert any(event.event_name == "artifact.verified" for event in list_task_events(task.id))
